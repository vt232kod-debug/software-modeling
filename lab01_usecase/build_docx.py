#!/usr/bin/env python3
"""Збирає report_content.md Лабораторної роботи №1 у .docx (і далі в .pdf).

Складальник — той самий, що в lab02_classes, з двома відмінностями:

  1) номер і назва роботи не прошиті в коді: титульний аркуш будується з
     розділу «## Титульний аркуш» самого звіту, тому тема в документі не
     може розійтися з темою в джерелі — це єдине джерело правди;
  2) розділ «Титульний аркуш» не дублюється в тілі документа: його блоки
     споживає титулка.

Запуск:  ~/.venvs/systems-ai/bin/python build_docx.py
"""
import re
import struct
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

HERE = Path(__file__).resolve().parent
MD = HERE / "report_content.md"
OUT = HERE / "МАПЗ-ЛР-1-ВТ-23-2-Камінський.docx"

TITLE_SECTION = "Титульний аркуш"

BODY_PT = 14          # основний текст
TABLE_PT = 11         # текст у таблицях
CODE_PT = 8.5         # блоки коду
MONO = "Courier New"

PORTRAIT = (Twips(12240), Twips(15840))      # Letter
MARGINS = dict(top=Twips(1134), bottom=Twips(1134),
               left=Twips(1417), right=Twips(850))

ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
         "left": WD_ALIGN_PARAGRAPH.LEFT}

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def png_size(path: Path):
    data = path.read_bytes()[:33]
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"{path} не PNG")
    return struct.unpack(">II", data[16:24])


def set_page(section, landscape: bool) -> None:
    w, h = PORTRAIT
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = h, w
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = w, h
    for k, v in MARGINS.items():
        setattr(section, f"{k}_margin", v)


def plain(text: str) -> str:
    """Знімає розмітку: у титулці вона не потрібна."""
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).replace("`", "")


def add_runs(par, text: str, size: float, bold: bool = False) -> None:
    """Розкладає текст на прогони, обробляючи **жирний** і `моноширинний`."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = par.add_run(piece[2:-2])
            r.bold = True
            r.font.size = Pt(size)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = MONO
            r.font.size = Pt(round(size * 0.9, 1))
            r.bold = bold
        else:
            r = par.add_run(piece)
            r.bold = bold
            r.font.size = Pt(size)


def add_paragraph(doc, text: str, size=BODY_PT, bold=False, align="both",
                  before=0, after=120) -> None:
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    p.paragraph_format.space_before = Twips(before)
    p.paragraph_format.space_after = Twips(after)
    add_runs(p, text, size, bold)


def add_code(doc, lines) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = ALIGN["left"]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = MONO
        r.font.size = Pt(CODE_PT)


CONTENT_W_IN = (12240 - 1417 - 850) / 1440      # ширина набору, портрет


def column_widths(rows, cols):
    """Ширини колонок: пропорційні обсягу тексту, але не вужчі за найдовше
    нерозривне слово в колонці (інакше назви акторів ламаються посередині)."""
    props, mins = [], []
    for j in range(cols):
        texts = [r[j] for r in rows if j < len(r)]
        longest = max((len(t) for t in texts), default=1)
        token = max((len(w.strip("`*")) for t in texts for w in t.split()),
                    default=1)
        props.append(max(min(longest, 140), 8) ** 0.75)
        mins.append(min(2.4, token * 0.082 + 0.22))
    total = sum(props)
    w = [max(CONTENT_W_IN * x / total, m) for x, m in zip(props, mins)]
    for _ in range(20):
        over = sum(w) - CONTENT_W_IN
        if over <= 1e-6:
            break
        slack = [wi - mi for wi, mi in zip(w, mins)]
        tot = sum(slack)
        if tot <= 1e-6:
            w = [wi * CONTENT_W_IN / sum(w) for wi in w]
            break
        w = [wi - over * sl / tot for wi, sl in zip(w, slack)]
    return w


def add_table(doc, rows) -> None:
    cols = len(rows[0])
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = column_widths(rows, cols)

    # фіксована розкладка + явні ширини в сітці, інакше LibreOffice і Word
    # перерозподіляють колонки порівну й довгі обґрунтування стискаються
    tbl_pr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    grid = t._tbl.find(qn("w:tblGrid"))
    for col_el, w in zip(grid.findall(qn("w:gridCol")), widths):
        col_el.set(qn("w:w"), str(int(w * 1440)))
    for i, row in enumerate(rows):
        tr = t.add_row()
        tr_pr = tr._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        cells = tr.cells
        for j in range(cols):
            cell = cells[j]
            cell.width = Inches(widths[j])
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = ALIGN["center"] if i == 0 else ALIGN["left"]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, row[j] if j < len(row) else "", TABLE_PT, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Twips(120)


def add_figure(doc, img: Path, caption: str) -> None:
    """Рисунок на сторінці в альбомній орієнтації (орієнтацією керує main)."""
    avail_w = 15840 - 1417 - 850          # twips
    avail_h = 12240 - 1134 - 1134
    max_w_in = avail_w / 1440
    max_h_in = avail_h / 1440 - 0.55      # місце під підпис

    w_px, h_px = png_size(img)
    scale = min(max_w_in / w_px, max_h_in / h_px)
    width_in = w_px * scale

    p = doc.add_paragraph()
    p.alignment = ALIGN["center"]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(img), width=Inches(width_in))

    cap = doc.add_paragraph()
    cap.alignment = ALIGN["center"]
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(0)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(TABLE_PT)

    # кегль підписів усередині діаграми після масштабування під аркуш:
    # PlantUML малює текст овалів 14 px, Mermaid і Draw.io — 2x-растр
    src_px = 14 if img.name.startswith("usecase") else 24
    pt = src_px / w_px * width_in * 72
    print(f"  {img.name}: {w_px}x{h_px} px -> {width_in:.2f} in, "
          f"кегль усередині ≈ {pt:.1f} pt")


def parse(md: str):
    """Проста потокова розбірка підмножини Markdown, яку вживає звіт."""
    out, lines, i = [], md.split("\n"), 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].startswith("```"):
                buf.append(lines[j]); j += 1
            out.append(("code", buf)); i = j + 1
        elif line.startswith("### "):
            out.append(("h2", line[4:].strip())); i += 1
        elif line.startswith("## "):
            out.append(("h1", line[3:].strip())); i += 1
        elif line.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            out.append(("img", (m.group(1), m.group(2)))); i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") and c for c in cells):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
        elif line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip()); i += 1
            out.append(("list", items))
        elif line.strip():
            out.append(("p", line.strip())); i += 1
        else:
            i += 1
    return out


def split_title(blocks):
    """Відокремлює розділ «Титульний аркуш» від тіла звіту.

    Повертає (поля титулки, решта блоків). Поля беруться з таблиці
    «Поле | Значення» цього розділу — у звіті вона єдине джерело правди
    про номер роботи, її назву й варіант предметної області.
    """
    if not blocks or blocks[0] != ("h1", TITLE_SECTION):
        raise SystemExit(f"звіт має починатися розділом «## {TITLE_SECTION}»")
    rest, fields, i = [], {}, 1
    while i < len(blocks) and blocks[i][0] != "h1":
        kind, payload = blocks[i]
        if kind == "table":
            for row in payload[1:]:
                fields[plain(row[0])] = plain(row[1])
        i += 1
    rest = blocks[i:]
    need = ["Міністерство", "Заклад освіти", "Кафедра", "Номер роботи",
            "Дисципліна", "Назва роботи", "Варіант предметної області",
            "Виконавець", "Група", "Місто й рік"]
    missing = [k for k in need if k not in fields]
    if missing:
        raise SystemExit("у титулці немає полів: " + ", ".join(missing))
    return fields, rest


def add_title_page(doc, f) -> None:
    rows = [
        (f["Міністерство"].upper(), "center", False, 14),
        (f["Заклад освіти"].upper(), "center", False, 14),
        (f["Кафедра"], "center", False, 14),
    ] + [("", "center", False, 14)] * 6 + [
        (f["Номер роботи"].upper(), "center", True, 16),
        (f"з дисципліни «{f['Дисципліна']}»", "center", False, 14),
        (f["Назва роботи"].upper(), "center", True, 14),
        ("", "center", False, 14),
        (f"Варіант предметної області: {f['Варіант предметної області']}",
         "center", False, 13),
    ] + [("", "center", False, 14)] * 5 + [
        (f"Виконав: студент групи {f['Група']}", "right", False, 14),
        (f["Виконавець"], "right", False, 14),
        ("Перевірив: ____________________", "right", False, 14),
    ] + [("", "center", False, 14)] * 5 + [
        (f["Місто й рік"], "center", False, 14),
    ]
    for text, align, bold, size in rows:
        p = doc.add_paragraph()
        p.alignment = ALIGN[align]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def main() -> None:
    doc = Document()
    set_page(doc.sections[0], landscape=False)
    style = doc.styles["Normal"]
    style.font.size = Pt(BODY_PT)
    style.paragraph_format.space_after = Twips(120)
    style.paragraph_format.line_spacing = 1.15

    fields, blocks = split_title(parse(MD.read_text(encoding="utf-8")))
    add_title_page(doc, fields)
    print(f"титулка: {fields['Номер роботи']} — {fields['Назва роботи']}")
    print(f"         варіант: {fields['Варіант предметної області'][:60]}…")

    figures = 0
    landscape = False
    for kind, payload in blocks:
        want_landscape = (kind == "img")
        if want_landscape != landscape:
            # перемикаємо орієнтацію лише там, де вона справді змінюється,
            # інакше між двома рисунками зʼявляється порожня сторінка
            set_page(doc.add_section(WD_SECTION.NEW_PAGE), want_landscape)
            landscape = want_landscape
        elif want_landscape:
            # кожен наступний рисунок — з нової сторінки тієї ж секції
            br = doc.add_paragraph()
            br.paragraph_format.space_after = Pt(0)
            br.add_run().add_break(WD_BREAK.PAGE)
        if kind == "h1":
            add_paragraph(doc, payload, BODY_PT, True, "left", 240, 120)
        elif kind == "h2":
            add_paragraph(doc, payload, BODY_PT - 1, True, "left", 180, 100)
        elif kind == "p":
            add_paragraph(doc, payload)
        elif kind == "list":
            for item in payload:
                p = doc.add_paragraph(style="List Bullet")
                p.alignment = ALIGN["both"]
                p.paragraph_format.space_after = Twips(60)
                add_runs(p, item, BODY_PT)
        elif kind == "code":
            add_code(doc, payload)
            doc.add_paragraph().paragraph_format.space_after = Twips(120)
        elif kind == "table":
            add_table(doc, payload)
        elif kind == "img":
            caption, fname = payload
            add_figure(doc, HERE / fname, caption)
            figures += 1

    doc.save(OUT)
    print(f"збережено {OUT.name}; рисунків: {figures}")

    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", str(HERE), str(OUT)],
                   check=True, capture_output=True)
    pdf = OUT.with_suffix(".pdf")
    print(f"збережено {pdf.name}: {pdf.stat().st_size // 1024} КБ")


if __name__ == "__main__":
    sys.exit(main())
